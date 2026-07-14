from __future__ import annotations

import json
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from post import (
    CHROME_USER_AGENT,
    VIEWPORT,
    RedditScreenshotError,
    browser_can_show_ui,
    dismiss_common_popups,
    hide_fixed_overlays,
    import_playwright,
    launch_browser,
)

OUTPUT_DIR = Path("instagram_shots")
EXCEL_PATH = Path("instagram.xlsx")
CONFIG_PATH = Path("instagramConfig.json")
DELAY_BETWEEN_JOBS_SECONDS = 3.0
CAPTURE_XPATH = "//section/main/div/div[1]/div"
MODAL_XPATH = '//*[@aria-modal="true" and @role="dialog"]'
CLOSE_MODAL_XPATH = "//section/main/div[1]/div[2]/div/div/div/div/div[1]/div"


@dataclass(frozen=True)
class InstagramJob:
    row_number: int
    stt: str
    channel: str
    link: str


@dataclass(frozen=True)
class InstagramConfig:
    start: int
    end: int | None


def log(message: str) -> None:
    print(f"[Instagram] {message}")


def cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def sanitize_output_name(value: str) -> str:
    stem = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._-") or "instagram"
    return stem


def load_jobs_from_excel(path: Path) -> list[InstagramJob]:
    if not path.exists():
        raise RedditScreenshotError(f"Khong tim thay file Excel: {path.resolve()}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        jobs: list[InstagramJob] = []
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            stt = cell_to_text(row[0]) if len(row) >= 1 else ""
            channel = cell_to_text(row[1]) if len(row) >= 2 else ""
            link = cell_to_text(row[2]) if len(row) >= 3 else ""

            if not any(cell_to_text(value) for value in row):
                continue
            if row_number == 1 and link.lower() == "link air":
                continue
            if not stt or not link:
                log(f"Bo qua dong {row_number}: thieu cot # hoac cot Link Air.")
                continue

            jobs.append(
                InstagramJob(
                    row_number=row_number,
                    stt=stt,
                    channel=channel,
                    link=link,
                )
            )
    finally:
        workbook.close()

    if not jobs:
        raise RedditScreenshotError("Khong co dong Instagram hop le nao trong Excel.")
    return jobs


def parse_config_int(config: dict[str, object], key: str) -> int | None:
    value = config.get(key)
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int):
        raise RedditScreenshotError(f"Gia tri '{key}' trong {CONFIG_PATH} phai la so nguyen.")
    if value < 1:
        raise RedditScreenshotError(f"Gia tri '{key}' trong {CONFIG_PATH} phai >= 1.")
    return value


def load_instagram_config(path: Path) -> InstagramConfig:
    if not path.exists():
        raise RedditScreenshotError(f"Khong tim thay file config: {path.resolve()}")

    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise RedditScreenshotError(f"File {path} khong phai JSON hop le: {exc}") from exc

    if not isinstance(config, dict):
        raise RedditScreenshotError(f"File {path} phai la mot JSON object.")

    start = parse_config_int(config, "start")
    end = parse_config_int(config, "end")
    if start is None:
        start = 1
    if end is not None and end < start:
        raise RedditScreenshotError("Gia tri 'end' phai lon hon hoac bang 'start'.")

    return InstagramConfig(start=start, end=end)


def filter_jobs_by_config(
    jobs: list[InstagramJob], config: InstagramConfig
) -> list[InstagramJob]:
    first_row = jobs[0].row_number
    last_row = jobs[-1].row_number
    effective_end = last_row if config.end is None else min(config.end, last_row)

    filtered = [job for job in jobs if config.start <= job.row_number <= effective_end]
    if not filtered:
        raise RedditScreenshotError(
            f"Khong co dong Instagram hop le nao trong khoang Excel row {config.start} den {effective_end}."
        )

    log(
        f"Chay Instagram tu dong Excel {config.start} den {effective_end} "
        f"(du lieu hop le tu {first_row} den {last_row})."
    )
    return filtered


def make_output_path(stt: str, url: str) -> Path:
    slug = sanitize_output_name(url.rstrip("/").split("/")[-1] or "instagram")
    stem = sanitize_output_name(stt)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    candidate = OUTPUT_DIR / f"#{stem}_{slug}.png"
    index = 2
    while candidate.exists():
        candidate = OUTPUT_DIR / f"#{stem}_{slug}_{index}.png"
        index += 1
    return candidate


def has_visible_modal(page) -> bool:
    modal = page.locator(f"xpath={MODAL_XPATH}")
    try:
        return modal.count() > 0 and modal.first.is_visible()
    except Exception:
        return False


def capture_target_is_inside_modal(page, capture_locator) -> bool:
    modal = page.locator(f"xpath={MODAL_XPATH}")
    if modal.count() == 0 or capture_locator.count() == 0:
        return False

    return bool(
        page.evaluate(
            """
            ({ modal, target }) => {
              if (!modal || !target) return false;
              return modal.contains(target);
            }
            """,
            {
                "modal": modal.first.element_handle(),
                "target": capture_locator.first.element_handle(),
            },
        )
    )


def click_locator_if_visible(locator, timeout: int = 1500) -> bool:
    try:
        if locator.count() > 0 and locator.first.is_visible():
            locator.first.click(timeout=timeout)
            return True
    except Exception:
        return False
    return False


def close_modal_if_needed(page, capture_locator) -> bool:
    if not has_visible_modal(page):
        # log("Khong thay modal aria-modal=true role=dialog.")
        return False

    if capture_target_is_inside_modal(page, capture_locator):
        # log("Target can chup nam ben trong modal, se giu modal de chup.")
        return False

    # log("Phat hien modal dang che noi dung, thu dong modal.")
    close_target = page.locator(f"xpath={CLOSE_MODAL_XPATH}")
    if click_locator_if_visible(close_target):
        page.wait_for_timeout(600)
        return True

    fallback_patterns = [
        re.compile(r"close", re.I),
        re.compile(r"not now", re.I),
        re.compile(r"dismiss", re.I),
    ]
    modal = page.locator(f"xpath={MODAL_XPATH}").first
    for pattern in fallback_patterns:
        try:
            button = modal.get_by_role("button", name=pattern)
            if click_locator_if_visible(button):
                page.wait_for_timeout(600)
                return True
        except Exception:
            continue

    try:
        page.keyboard.press("Escape")
        page.wait_for_timeout(400)
    except Exception:
        pass

    return not has_visible_modal(page)


def wait_for_capture_target(page):
    locator = page.locator(f"xpath={CAPTURE_XPATH}")
    locator.first.wait_for(state="visible", timeout=45_000)
    return locator.first


def prepare_capture(page, capture_locator) -> None:
    page.evaluate(
        """
        (node) => {
          if (!node) return;
          node.scrollIntoView({ block: "center", inline: "center", behavior: "instant" });
          document.documentElement.style.setProperty("scroll-behavior", "auto", "important");
          document.body.style.setProperty("background", "#ffffff", "important");
        }
        """,
        capture_locator.element_handle(),
    )
    hide_fixed_overlays(page, capture_locator.element_handle())
    page.wait_for_timeout(500)


def take_instagram_screenshot(url: str, output_path: Path, *, headless: bool) -> Path:
    sync_playwright, PlaywrightTimeoutError, _ = import_playwright()

    with sync_playwright() as playwright:
        browser = launch_browser(playwright, headless=headless)
        context = browser.new_context(
            viewport=VIEWPORT,
            device_scale_factor=2,
            color_scheme="light",
            locale="en-US",
            user_agent=CHROME_USER_AGENT,
            extra_http_headers={"Accept-Language": "en-US,en;q=0.9"},
        )
        page = context.new_page()
        page.set_default_timeout(45_000)
        page.add_init_script(
            """
            Object.defineProperty(navigator, 'webdriver', {
              get: () => undefined
            });
            """
        )

        try:
            # log(f"Mo trang: {url}")
            page.goto(url, wait_until="domcontentloaded")
            try:
                page.wait_for_load_state("networkidle", timeout=10_000)
            except PlaywrightTimeoutError:
                pass

            dismiss_common_popups(page)
            capture_locator = wait_for_capture_target(page)
            close_modal_if_needed(page, capture_locator)
            capture_locator = wait_for_capture_target(page)
            prepare_capture(page, capture_locator)

            # log(f"Chup target: {CAPTURE_XPATH}")
            capture_locator.screenshot(path=str(output_path))
            return output_path
        finally:
            context.close()
            browser.close()


def main() -> int:
    headless = True
    started = time.time()

    try:
        jobs = load_jobs_from_excel(EXCEL_PATH)
        config = load_instagram_config(CONFIG_PATH)
        filtered_jobs = filter_jobs_by_config(jobs, config)
    except RedditScreenshotError as exc:
        print(f"Loi: {exc}", file=sys.stderr)
        return 1

    for index, job in enumerate(filtered_jobs, start=1):
        output_path = make_output_path(job.stt, job.link).resolve()
        # log(
        #     f"[{index}/{len(filtered_jobs)}] Dong {job.row_number} | "
        #     f"#{job.stt} | {job.link}"
        # )
        try:
            result = take_instagram_screenshot(job.link, output_path, headless=headless)
            # log(f"Da luu anh tai: {result}")
        except Exception as exc:
            print(
                f"Loi o dong {job.row_number} ({job.link}): {exc}",
                file=sys.stderr,
            )
            return 1

        if index < len(filtered_jobs):
            time.sleep(DELAY_BETWEEN_JOBS_SECONDS)

    elapsed = time.time() - started
    # log(f"Hoan tat trong {elapsed:.1f}s")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
