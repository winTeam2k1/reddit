from __future__ import annotations

import json
import re
import sys
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from post import (
    CHROME_USER_AGENT,
    VIEWPORT,
    RedditScreenshotError,
    browser_can_show_ui,
    capture_post_slices,
    dismiss_common_popups,
    get_post_bounds,
    hide_fixed_overlays,
    import_playwright,
    launch_browser,
    sanitize_output_name,
    stitch_pngs,
)

OUTPUT_DIR = Path("facebook_shots")
EXCEL_PATH = Path("facebook.xlsx")
CONFIG_PATH = Path("facebookConfig.json")
DELAY_BETWEEN_JOBS_SECONDS = 3.0
MODAL_SELECTOR = '[aria-modal="true"][role="dialog"]'
CLOSE_SELECTOR = '[aria-label="Close"]'
CAPTURE_SELECTOR = '[role="feed"] > [role="presentation"]'
CAPTURE_FALLBACK_SELECTOR = '[role="feed"] [role="presentation"]'


@dataclass(frozen=True)
class FacebookJob:
    row_number: int
    stt: str
    channel: str
    link: str


@dataclass(frozen=True)
class FacebookConfig:
    start: int
    end: int | None


def log(message: str) -> None:
    print(f"[Facebook] {message}")


def cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_jobs_from_excel(path: Path) -> list[FacebookJob]:
    if not path.exists():
        raise RedditScreenshotError(f"Khong tim thay file Excel: {path.resolve()}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        jobs: list[FacebookJob] = []
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            stt = cell_to_text(row[0]) if len(row) >= 1 else ""
            channel = cell_to_text(row[1]) if len(row) >= 2 else ""
            link = cell_to_text(row[2]) if len(row) >= 3 else ""

            if not any(cell_to_text(value) for value in row):
                continue
            if row_number == 1 and link.lower() == "link air":
                continue
            if not stt or not link:
                log(f"Bo qua dong {row_number}: thieu cot # hoac cot Link Air.")
                continue

            jobs.append(
                FacebookJob(
                    row_number=row_number,
                    stt=stt,
                    channel=channel,
                    link=link,
                )
            )
    finally:
        workbook.close()

    if not jobs:
        raise RedditScreenshotError("Khong co dong Facebook hop le nao trong Excel.")
    return jobs


def parse_config_int(config: dict[str, object], key: str) -> int | None:
    value = config.get(key)
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int):
        raise RedditScreenshotError(f"Gia tri '{key}' trong {CONFIG_PATH} phai la so nguyen.")
    if value < 1:
        raise RedditScreenshotError(f"Gia tri '{key}' trong {CONFIG_PATH} phai >= 1.")
    return value


def load_facebook_config(path: Path) -> FacebookConfig:
    if not path.exists():
        raise RedditScreenshotError(f"Khong tim thay file config: {path.resolve()}")

    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise RedditScreenshotError(f"File {path} khong phai JSON hop le: {exc}") from exc

    if not isinstance(config, dict):
        raise RedditScreenshotError(f"File {path} phai la mot JSON object.")

    start = parse_config_int(config, "start")
    end = parse_config_int(config, "end")
    if start is None:
        start = 1
    if end is not None and end < start:
        raise RedditScreenshotError("Gia tri 'end' phai lon hon hoac bang 'start'.")

    return FacebookConfig(start=start, end=end)


def filter_jobs_by_config(
    jobs: list[FacebookJob], config: FacebookConfig
) -> list[FacebookJob]:
    first_row = jobs[0].row_number
    last_row = jobs[-1].row_number
    effective_end = last_row if config.end is None else min(config.end, last_row)

    filtered = [job for job in jobs if config.start <= job.row_number <= effective_end]
    if not filtered:
        raise RedditScreenshotError(
            f"Khong co dong Facebook hop le nao trong khoang Excel row {config.start} den {effective_end}."
        )

    log(
        f"Chay Facebook tu dong Excel {config.start} den {effective_end} "
        f"(du lieu hop le tu {first_row} den {last_row})."
    )
    return filtered


def make_output_path(stt: str, url: str) -> Path:
    slug = sanitize_output_name(url.rstrip("/").split("/")[-1] or "facebook")
    stem = sanitize_output_name(stt)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    candidate = OUTPUT_DIR / f"#{stem}_{slug}.png"
    index = 2
    while candidate.exists():
        candidate = OUTPUT_DIR / f"#{stem}_{slug}_{index}.png"
        index += 1
    return candidate


def has_visible_modal(page) -> bool:
    modal = page.locator(MODAL_SELECTOR)
    try:
        return modal.count() > 0 and modal.first.is_visible()
    except Exception:
        return False


def click_locator_if_visible(locator, timeout: int = 1500) -> bool:
    try:
        if locator.count() > 0 and locator.first.is_visible():
            locator.first.click(timeout=timeout)
            return True
    except Exception:
        return False
    return False


def close_modal_if_present(page) -> bool:
    if not has_visible_modal(page):
        return False

    modal = page.locator(MODAL_SELECTOR).first
    close_button = modal.locator(CLOSE_SELECTOR)
    if click_locator_if_visible(close_button):
        page.wait_for_timeout(800)
        return True

    role_button = modal.get_by_role("button", name=re.compile(r"close", re.I))
    if click_locator_if_visible(role_button):
        page.wait_for_timeout(800)
        return True

    page.keyboard.press("Escape")
    page.wait_for_timeout(500)
    return not has_visible_modal(page)


def wait_for_capture_target(page):
    primary = page.locator(CAPTURE_SELECTOR)
    try:
        primary.first.wait_for(state="visible", timeout=20_000)
        return primary.first
    except Exception:
        fallback = page.locator(CAPTURE_FALLBACK_SELECTOR)
        fallback.first.wait_for(state="visible", timeout=25_000)
        return fallback.first


def scroll_target_to_load_all_content(page, capture_locator) -> None:
    handle = capture_locator.element_handle()
    if handle is None:
        raise RedditScreenshotError("Khong lay duoc node cua target can chup.")

    page.evaluate(
        """
        async (node) => {
          if (!node) return;

          const sleep = (ms) => new Promise((resolve) => setTimeout(resolve, ms));
          node.scrollIntoView({ block: "center", inline: "center", behavior: "instant" });

          const canScroll = node.scrollHeight > node.clientHeight + 4;
          if (!canScroll) return;

          let lastHeight = 0;
          let stableRounds = 0;

          while (stableRounds < 2) {
            const currentHeight = node.scrollHeight;
            if (currentHeight === lastHeight) {
              stableRounds += 1;
            } else {
              stableRounds = 0;
              lastHeight = currentHeight;
            }

            for (let top = 0; top < node.scrollHeight; top += Math.max(400, node.clientHeight - 120)) {
              node.scrollTop = top;
              node.dispatchEvent(new Event("scroll", { bubbles: true }));
              await sleep(250);
            }

            node.scrollTop = node.scrollHeight;
            node.dispatchEvent(new Event("scroll", { bubbles: true }));
            await sleep(400);
          }

          node.scrollTop = 0;
          node.dispatchEvent(new Event("scroll", { bubbles: true }));
          await sleep(250);
        }
        """,
        handle,
    )


def expand_target_for_capture(page, capture_locator):
    handle = capture_locator.element_handle()
    if handle is None:
        raise RedditScreenshotError("Khong lay duoc node cua target can chup.")

    page.evaluate(
        """
        (node) => {
          if (!node) return;
          node.scrollIntoView({ block: "start", inline: "center", behavior: "instant" });
          const fullHeight = Math.max(
            node.scrollHeight,
            node.clientHeight,
            node.offsetHeight,
          );
          node.style.setProperty("overflow", "visible", "important");
          node.style.setProperty("max-height", "none", "important");
          node.style.setProperty("height", `${fullHeight}px`, "important");
          node.style.setProperty("contain", "none", "important");
          node.style.setProperty("display", "block", "important");
          node.style.setProperty("visibility", "visible", "important");
          node.style.setProperty("opacity", "1", "important");
          document.documentElement.style.setProperty("scroll-behavior", "auto", "important");
          document.body.style.setProperty("background", "#ffffff", "important");
        }
        """,
        handle,
    )
    hide_fixed_overlays(page, handle)
    page.wait_for_timeout(500)
    return handle


def save_capture_screenshot(page, capture_handle, output_path: Path) -> Path:
    bounds = get_post_bounds(capture_handle)
    with tempfile.TemporaryDirectory(prefix="facebook_capture_") as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        parts = capture_post_slices(page, bounds, temp_dir)
        if len(parts) == 1:
            output_path.write_bytes(parts[0].read_bytes())
        else:
            stitch_pngs(parts, output_path)
    return output_path


def ensure_not_facebook_gate(page) -> None:
    body_text = page.locator("body").inner_text(timeout=5_000)
    blocked_markers = [
        "Log in to continue",
        "You must log in to continue",
        "login to continue",
        "Something went wrong",
    ]
    for marker in blocked_markers:
        if marker.lower() in body_text.lower():
            raise RedditScreenshotError(f"Facebook dang chan hoac yeu cau dang nhap: '{marker}'.")


def take_facebook_screenshot_once(url: str, output_path: Path, *, headless: bool) -> Path:
    sync_playwright, PlaywrightTimeoutError, PlaywrightError = import_playwright()

    with sync_playwright() as playwright:
        browser = launch_browser(playwright, headless=headless)
        context = browser.new_context(
            viewport=VIEWPORT,
            device_scale_factor=2,
            color_scheme="light",
            locale="en-US",
            user_agent=CHROME_USER_AGENT,
            extra_http_headers={"Accept-Language": "en-US,en;q=0.9"},
        )
        page = context.new_page()
        page.set_default_timeout(45_000)
        page.add_init_script(
            """
            Object.defineProperty(navigator, 'webdriver', {
              get: () => undefined
            });
            """
        )

        try:
            page.goto(url, wait_until="domcontentloaded")
            try:
                page.wait_for_load_state("networkidle", timeout=12_000)
            except PlaywrightTimeoutError:
                pass

            dismiss_common_popups(page)
            close_modal_if_present(page)
            ensure_not_facebook_gate(page)

            capture_locator = wait_for_capture_target(page)
            scroll_target_to_load_all_content(page, capture_locator)
            capture_locator = wait_for_capture_target(page)
            capture_handle = expand_target_for_capture(page, capture_locator)
            return save_capture_screenshot(page, capture_handle, output_path)
        except PlaywrightTimeoutError as exc:
            raise RedditScreenshotError(
                f"Het thoi gian cho khi tai Facebook hoac tim {CAPTURE_SELECTOR}."
            ) from exc
        except PlaywrightError as exc:
            raise RedditScreenshotError(f"Playwright loi khi chup Facebook: {exc}") from exc
        finally:
            context.close()
            browser.close()


def take_facebook_screenshot(url: str, output_path: Path) -> Path:
    attempts = [True]
    if browser_can_show_ui():
        attempts.append(False)

    last_error: RedditScreenshotError | None = None
    for headless in attempts:
        try:
            return take_facebook_screenshot_once(url, output_path, headless=headless)
        except RedditScreenshotError as exc:
            last_error = exc
            if "dang chan hoac yeu cau dang nhap" not in str(exc) or not headless:
                raise

    if last_error is None:
        raise RedditScreenshotError("Khong the chup bai Facebook.")
    raise RedditScreenshotError(
        f"{last_error} Thu dang nhap Facebook hoac chay browser hien hinh."
    )


def capture_jobs(jobs: list[FacebookJob]) -> int:
    failures = 0
    total = len(jobs)

    for index, job in enumerate(jobs, start=1):
        try:
            log(
                f"[{index}/{total}] Dong {job.row_number} | "
                f"#{job.stt} | {job.link}"
            )
            output_path = make_output_path(job.stt, job.link)
            final_path = take_facebook_screenshot(job.link, output_path)
            log(f"Luu anh tai: {final_path.resolve()}")
        except RedditScreenshotError as exc:
            failures += 1
            log(f"[{index}/{total}] Dong {job.row_number} loi: {exc}")

        if index < total and DELAY_BETWEEN_JOBS_SECONDS > 0:
            time.sleep(DELAY_BETWEEN_JOBS_SECONDS)

    return failures


def main() -> int:
    try:
        jobs = load_jobs_from_excel(EXCEL_PATH)
        config = load_facebook_config(CONFIG_PATH)
        jobs = filter_jobs_by_config(jobs, config)
        failures = capture_jobs(jobs)
        if failures:
            log(f"Hoan tat voi {failures} dong loi.")
            return 1
        log("Hoan tat tat ca anh Facebook.")
        return 0
    except KeyboardInterrupt:
        log("Da huy.")
        return 130
    except RedditScreenshotError as exc:
        print(f"Loi: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
