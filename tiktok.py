from __future__ import annotations

import json
import re
import time
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook

from post import (
    CHROME_USER_AGENT,
    VIEWPORT,
    RedditScreenshotError,
    dismiss_common_popups,
    hide_fixed_overlays,
    import_playwright,
    launch_browser,
    sanitize_output_name,
)

OUTPUT_DIR = Path("tiktok_shots")
EXCEL_PATH = Path("tiktok.xlsx")
CONFIG_PATH = Path("tiktokConfig.json")
DELAY_BETWEEN_JOBS_SECONDS = 3.0


@dataclass(frozen=True)
class TikTokJob:
    row_number: int
    stt: str
    channel: str
    link: str
    video_id: str


@dataclass(frozen=True)
class TikTokConfig:
    start: int
    end: int | None


def log(message: str) -> None:
    print(f"[TikTok] {message}")


def cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_tiktok_video_id(url: str) -> str:
    path = urlparse(url).path
    match = re.search(r"/video/(\d+)", path)
    if not match:
        raise RedditScreenshotError(f"Link TikTok khong chua video id hop le: {url}")
    return match.group(1)


def load_jobs_from_excel(path: Path) -> list[TikTokJob]:
    if not path.exists():
        raise RedditScreenshotError(f"Khong tim thay file Excel: {path.resolve()}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        jobs: list[TikTokJob] = []
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            stt = cell_to_text(row[0]) if len(row) >= 1 else ""
            channel = cell_to_text(row[1]) if len(row) >= 2 else ""
            link = cell_to_text(row[2]) if len(row) >= 3 else ""

            if not any(cell_to_text(value) for value in row):
                continue
            if row_number == 1 and link.lower() == "link air":
                continue
            if not stt or not link:
                log(f"Bo qua dong {row_number}: thieu cot # hoac cot Link Air.")
                continue

            jobs.append(
                TikTokJob(
                    row_number=row_number,
                    stt=stt,
                    channel=channel,
                    link=link,
                    video_id=parse_tiktok_video_id(link),
                )
            )
    finally:
        workbook.close()

    if not jobs:
        raise RedditScreenshotError("Khong co dong TikTok hop le nao trong Excel.")
    return jobs


def parse_config_int(config: dict[str, object], key: str) -> int | None:
    value = config.get(key)
    if value is None:
        return None
    if isinstance(value, bool) or not isinstance(value, int):
        raise RedditScreenshotError(f"Gia tri '{key}' trong {CONFIG_PATH} phai la so nguyen.")
    if value < 1:
        raise RedditScreenshotError(f"Gia tri '{key}' trong {CONFIG_PATH} phai >= 1.")
    return value


def load_tiktok_config(path: Path) -> TikTokConfig:
    if not path.exists():
        raise RedditScreenshotError(f"Khong tim thay file config: {path.resolve()}")

    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise RedditScreenshotError(f"File {path} khong phai JSON hop le: {exc}") from exc

    if not isinstance(config, dict):
        raise RedditScreenshotError(f"File {path} phai la mot JSON object.")

    start = parse_config_int(config, "start")
    end = parse_config_int(config, "end")
    if start is None:
        start = 1
    if end is not None and end < start:
        raise RedditScreenshotError("Gia tri 'end' phai lon hon hoac bang 'start'.")

    return TikTokConfig(start=start, end=end)


def filter_jobs_by_config(jobs: list[TikTokJob], config: TikTokConfig) -> list[TikTokJob]:
    first_row = jobs[0].row_number
    last_row = jobs[-1].row_number
    effective_end = last_row if config.end is None else min(config.end, last_row)

    filtered = [job for job in jobs if config.start <= job.row_number <= effective_end]
    if not filtered:
        raise RedditScreenshotError(
            f"Khong co dong TikTok hop le nao trong khoang Excel row {config.start} den {effective_end}."
        )

    log(
        f"Chay TikTok tu dong Excel {config.start} den {effective_end} "
        f"(du lieu hop le tu {first_row} den {last_row})."
    )
    return filtered


def make_output_path(stt: str, video_id: str) -> Path:
    stem = sanitize_output_name(stt)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    candidate = OUTPUT_DIR / f"#{stem}_{video_id}.png"
    index = 2
    while candidate.exists():
        candidate = OUTPUT_DIR / f"#{stem}_{video_id}_{index}.png"
        index += 1
    return candidate


def ensure_not_tiktok_blocked(page) -> None:
    body_text = page.locator("body").inner_text(timeout=5_000)
    blocked_markers = [
        "Something went wrong",
        "maximum number of attempts",
        "Verify to continue",
        "Access denied",
    ]
    for marker in blocked_markers:
        if marker in body_text:
            raise RedditScreenshotError(f"TikTok dang chan hoac yeu cau xac minh: '{marker}'.")


def find_capture_container(page, video_id: str):
    wrapper_id = f"xgwrapper-0-{video_id}"
    handle = page.evaluate_handle(
        """
        (wrapperId) => {
          const wrapper = document.getElementById(wrapperId);
          if (!wrapper) return null;

          const article = wrapper.closest('article');
          const mediaCard = wrapper.closest('section[data-e2e="feed-video"]');
          const playerWrapper = wrapper.closest('[class*="DivBasicPlayerWrapper"]');

          const preferred =
            article ||
            mediaCard ||
            playerWrapper ||
            wrapper.parentElement;

          return preferred || wrapper;
        }
        """,
        wrapper_id,
    )
    return handle.as_element()


def wait_for_capture_container(page, video_id: str):
    wrapper_selector = f"#{'xgwrapper-0-' + video_id}"
    log(f"Cho xuat hien {wrapper_selector}")
    page.wait_for_selector(wrapper_selector, state="attached", timeout=45_000)

    for _ in range(30):
        element = find_capture_container(page, video_id)
        if element is not None:
            return element
        page.wait_for_timeout(500)

    raise RedditScreenshotError(
        f"Khong tim thay container cha cua phan tu xgwrapper-0-{video_id}."
    )


def prepare_page_for_capture(page, container_handle) -> None:
    page.evaluate(
        """
        (container) => {
          if (!container) return;
          const article = container.closest('article') || container;
          const actionBar = article.querySelector('section[class*="SectionActionBarContainer"]');
          const nestedWrapper = container.querySelector('[id^="xgwrapper-0-"]');
          if (nestedWrapper) {
            nestedWrapper.style.setProperty("display", "block", "important");
            nestedWrapper.style.setProperty("visibility", "visible", "important");
            nestedWrapper.style.setProperty("opacity", "1", "important");
          }

          article.style.setProperty("display", "block", "important");
          article.style.setProperty("visibility", "visible", "important");
          article.style.setProperty("opacity", "1", "important");
          article.style.setProperty("overflow", "visible", "important");
          article.style.setProperty("height", "auto", "important");

          container.style.setProperty("display", "block", "important");
          container.style.setProperty("visibility", "visible", "important");
          container.style.setProperty("opacity", "1", "important");
          container.style.setProperty("overflow", "visible", "important");
          container.style.setProperty("height", "auto", "important");

          if (actionBar) {
            actionBar.style.setProperty("display", "flex", "important");
            actionBar.style.setProperty("visibility", "visible", "important");
            actionBar.style.setProperty("opacity", "1", "important");
            actionBar.style.setProperty("overflow", "visible", "important");
            actionBar.style.setProperty("height", "auto", "important");
          }

          document.body.style.setProperty("background", "#000000", "important");
          document.documentElement.style.setProperty("background", "#000000", "important");
          article.scrollIntoView({ block: "center", inline: "center", behavior: "instant" });
        }
        """,
        container_handle,
    )
    hide_fixed_overlays(page, container_handle)
    page.wait_for_timeout(500)


def wait_for_action_bar(page, container_handle) -> None:
    try:
        page.wait_for_function(
            """
            (container) => {
              const article = container?.closest('article') || container;
              if (!article) return false;
              const actionBar = article.querySelector('section[class*="SectionActionBarContainer"]');
              if (!actionBar) return false;
              const buttons = actionBar.querySelectorAll('button');
              return buttons.length >= 4;
            }
            """,
            container_handle,
            timeout=15_000,
        )
    except Exception:
        pass


def save_container_screenshot(container_handle, output_path: Path) -> Path:
    container_handle.screenshot(path=str(output_path))
    return output_path


def take_tiktok_screenshot_once(job: TikTokJob, output_path: Path, *, headless: bool) -> Path:
    sync_playwright, PlaywrightTimeoutError, PlaywrightError = import_playwright()

    with sync_playwright() as playwright:
        browser = launch_browser(playwright, headless=headless)
        context = browser.new_context(
            viewport=VIEWPORT,
            device_scale_factor=2,
            color_scheme="dark",
            locale="en-US",
            user_agent=CHROME_USER_AGENT,
            extra_http_headers={"Accept-Language": "en-US,en;q=0.9"},
        )
        page = context.new_page()
        page.set_default_timeout(45_000)
        page.add_init_script(
            """
            Object.defineProperty(navigator, 'webdriver', {
              get: () => undefined
            });
            """
        )

        try:
            log(f"Mo link: {job.link}")
            page.goto(job.link, wait_until="domcontentloaded")
            try:
                page.wait_for_load_state("networkidle", timeout=12_000)
            except PlaywrightTimeoutError:
                pass

            dismiss_common_popups(page)
            ensure_not_tiktok_blocked(page)

            container_handle = wait_for_capture_container(page, job.video_id)
            wait_for_action_bar(page, container_handle)
            prepare_page_for_capture(page, container_handle)
            final_path = save_container_screenshot(container_handle, output_path)
            log(f"Da chup xong: {final_path.resolve()}")
            return final_path
        except PlaywrightTimeoutError as exc:
            raise RedditScreenshotError(
                f"Het thoi gian cho khi tai TikTok hoac tim xgwrapper-0-{job.video_id}."
            ) from exc
        except PlaywrightError as exc:
            raise RedditScreenshotError(f"Playwright loi khi chup TikTok: {exc}") from exc
        finally:
            context.close()
            browser.close()


def take_tiktok_screenshot(job: TikTokJob, output_path: Path) -> Path:
    return take_tiktok_screenshot_once(job, output_path, headless=True)


def capture_jobs(jobs: list[TikTokJob]) -> int:
    failures = 0
    total = len(jobs)

    for index, job in enumerate(jobs, start=1):
        try:
            log(
                f"[{index}/{total}] STT {job.stt} | Channel: {job.channel or 'TikTok'} | Video ID: {job.video_id}"
            )
            output_path = make_output_path(job.stt, job.video_id)
            final_path = take_tiktok_screenshot(job, output_path)
            log(f"Luu anh tai: {final_path.resolve()}")
        except RedditScreenshotError as exc:
            failures += 1
            log(f"[{index}/{total}] STT {job.stt} - loi tai dong {job.row_number}: {exc}")

        if index < total and DELAY_BETWEEN_JOBS_SECONDS > 0:
            time.sleep(DELAY_BETWEEN_JOBS_SECONDS)

    return failures


def main() -> int:
    try:
        jobs = load_jobs_from_excel(EXCEL_PATH)
        config = load_tiktok_config(CONFIG_PATH)
        jobs = filter_jobs_by_config(jobs, config)
        log(f"Tim thay {len(jobs)} link hop le trong {EXCEL_PATH.resolve()}")
        failures = capture_jobs(jobs)
        if failures:
            log(f"Hoan tat voi {failures} dong loi.")
            return 1
        log("Hoan tat tat ca anh TikTok.")
        return 0
    except KeyboardInterrupt:
        log("Da huy.")
        return 130
    except RedditScreenshotError as exc:
        log(f"Loi: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
